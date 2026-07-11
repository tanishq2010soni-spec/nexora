from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
from typing import Any

from nexora_ai.domain.entities.tool import ToolResult
from nexora_ai.domain.enums.tool_enums import ToolCategory, ToolPermission
from nexora_ai.domain.interfaces.logging_interface import LoggingInterface
from nexora_ai.infrastructure.tools.tool_registry import ToolRegistry

from backend.services.permissions_manager import PermissionsManager


class FileIntelligence:
    category = ToolCategory.FILESYSTEM
    permissions = [ToolPermission.READ]
    description = "File reading and intelligence tools"

    def __init__(
        self,
        tool_registry: ToolRegistry,
        permissions_manager: PermissionsManager,
        logger: LoggingInterface,
    ) -> None:
        self._tool_registry = tool_registry
        self._permissions = permissions_manager
        self._logger = logger

    async def register_tools(self) -> None:
        tools: list[tuple[str, Any]] = [
            ("read_file", self._make_handler(self._read_file, "Read a text file")),
            ("read_pdf", self._make_handler(self._read_pdf, "Read a PDF file")),
            ("read_docx", self._make_handler(self._read_docx, "Read a DOCX file")),
            ("read_excel", self._make_handler(self._read_excel, "Read an Excel file")),
            ("read_csv", self._make_handler(self._read_csv, "Read a CSV file")),
            ("read_markdown", self._make_handler(self._read_markdown, "Read a Markdown file")),
            ("read_json", self._make_handler(self._read_json, "Read a JSON file")),
            ("list_directory", self._make_handler(self._list_directory, "List directory contents")),
            ("get_file_info", self._make_handler(self._get_file_info, "Get file metadata")),
            ("search_files", self._make_handler(self._search_files, "Search files by name/pattern")),
        ]
        for name, handler in tools:
            await self._tool_registry.register_tool(name, handler)

    def _make_handler(self, coro_func: Any, desc: str) -> Any:
        handler = type(
            f"Tool_{coro_func.__name__}",
            (),
            {
                "category": ToolCategory.FILESYSTEM,
                "description": desc,
                "permissions": [ToolPermission.READ],
                "parameters": {},
                "execute": lambda params, cf=coro_func: cf(params),
            },
        )()
        return handler

    async def _read_file(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_file")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_file")
        try:
            content = Path(path).read_text(encoding="utf-8")
            return ToolResult(success=True, output=content, tool_name="read_file")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_file")

    async def _read_pdf(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_pdf")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_pdf")
        try:
            import fitz
            doc = fitz.open(path)
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return ToolResult(success=True, output=text, tool_name="read_pdf")
        except ImportError:
            try:
                content = Path(path).read_text(encoding="utf-8", errors="ignore")
                return ToolResult(success=True, output=f"[PDF raw text fallback]\n{content[:10000]}", tool_name="read_pdf")
            except Exception as exc:
                return ToolResult(success=False, error=f"PyMuPDF not available: {exc}", tool_name="read_pdf")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_pdf")

    async def _read_docx(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_docx")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_docx")
        try:
            import docx
            doc = docx.Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return ToolResult(success=True, output=text, tool_name="read_docx")
        except ImportError:
            return ToolResult(success=False, error="python-docx not available", tool_name="read_docx")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_docx")

    async def _read_excel(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_excel")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_excel")
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            result: dict[str, list[list[str]]] = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[list[str]] = []
                for row in ws.iter_row():
                    rows.append([str(cell.value or "") for cell in row])
                result[sheet_name] = rows
            wb.close()
            return ToolResult(success=True, output=result, tool_name="read_excel")
        except ImportError:
            return ToolResult(success=False, error="openpyxl not available", tool_name="read_excel")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_excel")

    async def _read_csv(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_csv")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_csv")
        try:
            with open(path, newline="", encoding="utf-8") as f:
                reader = csv.reader(f)
                rows = [row for row in reader]
            return ToolResult(success=True, output=rows, tool_name="read_csv")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_csv")

    async def _read_markdown(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_markdown")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_markdown")
        try:
            content = Path(path).read_text(encoding="utf-8")
            return ToolResult(success=True, output=content, tool_name="read_markdown")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_markdown")

    async def _read_json(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="read_json")
        req = self._permissions.request_permission("read_file", {"resource": path})
        if req.status == "denied":
            return ToolResult(success=False, error="Permission denied", tool_name="read_json")
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
            return ToolResult(success=True, output=data, tool_name="read_json")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="read_json")

    async def _list_directory(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", ".")
        try:
            p = Path(path)
            if not p.exists():
                return ToolResult(success=False, error=f"Path not found: {path}", tool_name="list_directory")
            entries: list[dict[str, Any]] = []
            for entry in sorted(p.iterdir()):
                try:
                    stat = entry.stat()
                    entries.append({
                        "name": entry.name,
                        "path": str(entry.absolute()),
                        "is_dir": entry.is_dir(),
                        "size": stat.st_size,
                        "modified": stat.st_mtime,
                    })
                except Exception:
                    entries.append({"name": entry.name, "path": str(entry.absolute()), "is_dir": entry.is_dir()})
            return ToolResult(success=True, output=entries, tool_name="list_directory")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="list_directory")

    async def _get_file_info(self, parameters: dict[str, Any]) -> ToolResult:
        path = parameters.get("path", "")
        if not path:
            return ToolResult(success=False, error="path is required", tool_name="get_file_info")
        try:
            p = Path(path)
            if not p.exists():
                return ToolResult(success=False, error=f"File not found: {path}", tool_name="get_file_info")
            stat = p.stat()
            info: dict[str, Any] = {
                "name": p.name,
                "path": str(p.absolute()),
                "size": stat.st_size,
                "is_dir": p.is_dir(),
                "is_file": p.is_file(),
                "extension": p.suffix,
                "created": stat.st_ctime,
                "modified": stat.st_mtime,
                "accessed": stat.st_atime,
            }
            return ToolResult(success=True, output=info, tool_name="get_file_info")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="get_file_info")

    async def _search_files(self, parameters: dict[str, Any]) -> ToolResult:
        query = parameters.get("query", "")
        path = parameters.get("path", ".")
        if not query:
            return ToolResult(success=False, error="query is required", tool_name="search_files")
        try:
            root = Path(path)
            if not root.exists():
                return ToolResult(success=False, error=f"Path not found: {path}", tool_name="search_files")
            results: list[dict[str, Any]] = []
            for entry in root.rglob("*"):
                if query.lower() in entry.name.lower():
                    try:
                        stat = entry.stat()
                        results.append({
                            "name": entry.name,
                            "path": str(entry.absolute()),
                            "is_dir": entry.is_dir(),
                            "size": stat.st_size,
                        })
                    except Exception:
                        results.append({"name": entry.name, "path": str(entry.absolute()), "is_dir": entry.is_dir()})
            return ToolResult(success=True, output=results[:200], tool_name="search_files")
        except Exception as exc:
            return ToolResult(success=False, error=str(exc), tool_name="search_files")
